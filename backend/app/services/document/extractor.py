import csv
import io
import json
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".csv",
    ".json",
}


def clean_text(text: str) -> str:
    lines = []

    for line in text.splitlines():
        line = " ".join(line.split())

        if line:
            lines.append(line)

    return "\n".join(lines)


def extract_pdf(file_path: Path) -> str:
    reader = PdfReader(str(file_path))

    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        if text.strip():
            pages.append(
                f"[Page {page_number}]\n{text}"
            )

    return "\n\n".join(pages)


def extract_docx(file_path: Path) -> str:
    document = DocxDocument(str(file_path))

    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text)

    for table_index, table in enumerate(
        document.tables,
        start=1,
    ):
        paragraphs.append(
            f"[Table {table_index}]"
        )

        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            paragraphs.append(
                " | ".join(values)
            )

    return "\n".join(paragraphs)


def extract_text_file(file_path: Path) -> str:
    return file_path.read_text(
        encoding="utf-8",
        errors="ignore",
    )


def extract_csv(file_path: Path) -> str:
    rows = []

    with file_path.open(
        "r",
        encoding="utf-8",
        errors="ignore",
        newline="",
    ) as file:

        reader = csv.reader(file)

        for row in reader:
            rows.append(" | ".join(row))

    return "\n".join(rows)


def extract_json(file_path: Path) -> str:
    raw = file_path.read_text(
        encoding="utf-8",
        errors="ignore",
    )

    try:
        data = json.loads(raw)

        return json.dumps(
            data,
            indent=2,
            ensure_ascii=False,
        )

    except json.JSONDecodeError:
        return raw


def extract_text(
    file_path: Path,
) -> str:

    extension = file_path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    if extension == ".pdf":
        text = extract_pdf(file_path)

    elif extension == ".docx":
        text = extract_docx(file_path)

    elif extension in {".txt", ".md"}:
        text = extract_text_file(file_path)

    elif extension == ".csv":
        text = extract_csv(file_path)

    elif extension == ".json":
        text = extract_json(file_path)

    else:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    cleaned = clean_text(text)

    if not cleaned:
        raise ValueError(
            "No readable text was found in the document."
        )

    return cleaned